"""excel-workbook / inspect_workbook — read-only inventory of an XLSX.

Parameter resolution order (dcc-mcp-core execute_script convention):
1. stdin JSON: {"input": ...}
2. CLI flags: --input
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def _force_utf8_stdio() -> None:
    """Deterministic output contract: stdout/stderr are always UTF-8."""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def run(params: dict) -> None:
    path = Path(params["input"])
    if not path.is_file():
        print(json.dumps({"success": False, "message": f"workbook not found: {path}", "context": {}}, ensure_ascii=False))
        return
    workbook = load_workbook(filename=str(path), data_only=False)
    try:
        sheets = []
        total_cells = 0
        total_formulas = 0
        for title in workbook.sheetnames:
            sheet = workbook[title]
            cells = 0
            formulas = 0
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cells += 1
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas += 1
            total_cells += cells
            total_formulas += formulas
            sheets.append(
                {
                    "name": title,
                    "dimensions": sheet.dimensions,
                    "populated_cells": cells,
                    "formulas": formulas,
                    "data_validations": len(sheet.data_validations.dataValidation),
                    "conditional_format_ranges": len(sheet.conditional_formatting._cf_rules),
                }
            )
        print(
            json.dumps(
                {
                    "success": True,
                    "message": f"inspected {path.name}: {len(sheets)} sheet(s)",
                    "context": {
                        "path": str(path),
                        "sheet_count": len(sheets),
                        "populated_cells": total_cells,
                        "formulas": total_formulas,
                        "defined_names": list(workbook.defined_names),
                        "sheets": sheets,
                    },
                },
                ensure_ascii=False,
            )
        )
    finally:
        workbook.close()


def main() -> None:
    _force_utf8_stdio()
    params: dict = {}
    if not sys.stdin.isatty():
        raw = sys.stdin.read()
        if raw.strip():
            try:
                params = json.loads(raw)
            except json.JSONDecodeError:
                params = {}
    if not params:
        parser = argparse.ArgumentParser(description="Read-only inventory of an XLSX workbook")
        parser.add_argument("--input", required=True, help="XLSX file path")
        params = vars(parser.parse_args())
    try:
        run(params)
    except Exception as exc:  # noqa: BLE001 — surface as structured error
        print(json.dumps({"success": False, "message": str(exc), "context": {}}, ensure_ascii=False))
        sys.exit(1)


if __name__ == "__main__":
    main()
