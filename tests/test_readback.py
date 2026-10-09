"""Read-back verification tests — the 1.0 write-then-read-back gate."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from dcc_mcp_excel.compiler import compile_workbook
from dcc_mcp_excel.readback import read_back
from dcc_mcp_excel.workbook_io import cell_values_match, table_display_name
from dcc_mcp_excel.workbook_ir import (
    Metadata,
    WorkbookEnvelope,
    WorkbookIr,
    Worksheet,
    load_workbook_ir,
    parse_envelope,
)

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"


def _envelope(document: WorkbookIr, document_id: str = "draft:test") -> WorkbookEnvelope:
    return WorkbookEnvelope(
        schema_version="office-ir/1.0",
        kind="workbook",
        document_id=document_id,
        metadata=Metadata(title="Test"),
        document=document,
    )


def test_read_back_matches_every_cell(tmp_path: Path) -> None:
    envelope = _envelope(
        WorkbookIr(
            worksheets=(
                Worksheet(name="Shots", rows=(("shot", "status", "frames"), ("sh010", "ip", 120))),
                Worksheet(name="Summary", rows=(("total", 1),)),
            )
        )
    )
    out = compile_workbook(envelope, tmp_path / "rb.xlsx")
    report = read_back(envelope, out)
    assert report.ok, report.to_dict()
    # Every positioned cell in the IR is checked: 3x2 on Shots + 1x2 on Summary.
    assert report.checked_cells == 8
    assert report.sheets == ("Shots", "Summary")


def test_read_back_detects_mismatches(tmp_path: Path) -> None:
    """The gate must fail, not warn, when the artifact does not match the IR."""
    envelope = _envelope(WorkbookIr(worksheets=(Worksheet(name="S", rows=(("expected",),)),)))
    out = compile_workbook(envelope, tmp_path / "rb.xlsx")

    # Same artifact, different IR: one cell disagrees.
    altered = _envelope(WorkbookIr(worksheets=(Worksheet(name="S", rows=(("changed",),)),)))
    report = read_back(altered, out)
    assert not report.ok
    assert report.mismatches[0].address == "A1"
    assert report.mismatches[0].expected == "changed"
    assert report.mismatches[0].actual == "expected"


def test_read_back_rejects_sheet_count_mismatch(tmp_path: Path) -> None:
    envelope = _envelope(WorkbookIr(worksheets=(Worksheet(name="A", rows=(("x",),)),)))
    out = compile_workbook(envelope, tmp_path / "rb.xlsx")
    wider = _envelope(
        WorkbookIr(
            worksheets=(Worksheet(name="A", rows=(("x",),)), Worksheet(name="B", rows=(("y",),))),
        )
    )
    with pytest.raises(ValueError, match="sheet count mismatch"):
        read_back(wider, out)


def test_read_back_rejects_title_mismatch(tmp_path: Path) -> None:
    envelope = _envelope(WorkbookIr(worksheets=(Worksheet(name="Alpha", rows=(("x",),)),)))
    out = compile_workbook(envelope, tmp_path / "rb.xlsx")
    renamed = _envelope(WorkbookIr(worksheets=(Worksheet(name="Beta", rows=(("x",),)),)))
    with pytest.raises(ValueError, match="expected title"):
        read_back(renamed, out)


def test_read_back_detects_formulas_present(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:f",
            "metadata": {"title": "f"},
            "document": {
                "worksheets": [{"name": "S", "rows": [["a", 2]]}],
                "formulas": [{"worksheet": "S", "cell": "C1", "formula": "=B1*2"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "f.xlsx")
    report = read_back(envelope, out)
    assert report.ok
    # The formula text survives; its value does not (evaluation is host_limited).
    assert report.formulas_present == 1


def test_read_back_reports_features(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:feat",
            "metadata": {"title": "feat"},
            "document": {
                "worksheets": [{"name": "S", "rows": [["n", 3]]}],
                "named_ranges": [{"name": "Total", "refers_to": "S!$A$1:$B$1"}],
                "validations": [{"worksheet": "S", "range": "A2:A9", "kind": "list", "params": {"values": ["a", "b"]}}],
                "conditional_formats": [
                    {"worksheet": "S", "range": "B2:B9", "kind": "cell_value", "params": {"formula": ["5"]}}
                ],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "feat.xlsx")
    report = read_back(envelope, out)
    assert report.ok
    assert report.named_ranges == ("Total",)
    assert report.validations == 1
    assert report.conditional_formats == 1


def test_read_back_checks_feature_on_sanitized_sheet_title(tmp_path: Path) -> None:
    """The gate must hold when the sheet title needed sanitizing.

    `Shot/List` compiles to `Shot-List`; every feature still has to be found
    on that sheet. This is the combination the old suite never covered.
    """
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:sanitized",
            "metadata": {"title": "sanitized"},
            "document": {
                "worksheets": [{"name": "Shot/List", "rows": [["n", "v"], ["a", "3"]]}],
                "named_ranges": [{"name": "Total", "refers_to": "'Shot/List'!$A$1:$B$2"}],
                "formulas": [{"worksheet": "Shot/List", "cell": "C1", "formula": "=B1*2"}],
                "validations": [
                    {"worksheet": "Shot/List", "range": "A2:A9", "kind": "list", "params": {"values": ["a", "b"]}}
                ],
                "conditional_formats": [
                    {"worksheet": "Shot/List", "range": "B2:B9", "kind": "cell_value", "params": {"formula": ["5"]}}
                ],
                "tables": [{"worksheet": "Shot/List", "range": "A1:B2", "name": "ShotTable"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "sanitized.xlsx")
    report = read_back(envelope, out)
    assert report.ok, report.to_dict()
    assert report.formulas_present == 1
    assert report.validations == 1
    assert report.conditional_formats == 1


def test_read_back_accepts_unnamed_table(tmp_path: Path) -> None:
    """`TableSpec.name` is optional, so an unnamed table must not fail the gate.

    The compiler derives a display name from the range; the gate has to look
    for that same derived name, or a correctly written table reads as missing.
    """
    ir = {
        "schema_version": "office-ir/1.0",
        "kind": "workbook",
        "document_id": "draft:unnamed",
        "metadata": {"title": "unnamed"},
        "document": {
            "worksheets": [{"name": "Shots", "rows": [["shot", "status"], ["sh010", "ip"]]}],
            "tables": [{"worksheet": "Shots", "range": "A1:B2"}],
        },
    }
    envelope = load_workbook_ir(ir)
    out = compile_workbook(envelope, tmp_path / "unnamed.xlsx")
    assert list(load_workbook(str(out))["Shots"].tables) == ["TableA1_B2"]
    report = read_back(envelope, out)
    assert report.ok, report.to_dict()
    assert report.feature_mismatches == ()


def test_unnamed_table_display_name_matches_compiler(tmp_path: Path) -> None:
    """Both sides derive the fallback name through one shared helper."""
    ir = {
        "schema_version": "office-ir/1.0",
        "kind": "workbook",
        "document_id": "draft:anchored",
        "metadata": {"title": "anchored"},
        "document": {
            "worksheets": [{"name": "S", "rows": [["a", "b"], ["c", "d"]]}],
            "tables": [{"worksheet": "S", "range": "$A$1:$B$2"}],
        },
    }
    envelope = load_workbook_ir(ir)
    out = compile_workbook(envelope, tmp_path / "anchored.xlsx")
    assert table_display_name(envelope.document.tables[0]) == "TableA1_B2"
    assert list(load_workbook(str(out))["S"].tables) == ["TableA1_B2"]
    assert read_back(envelope, out).ok


def test_read_back_detects_dropped_features(tmp_path: Path) -> None:
    """The gate must fail when a declared feature is missing from the artifact.

    Same artifact, richer IR: the extra features were never compiled, so a
    gate that only counts them would stay green. This pins the comparison.
    """
    base = {
        "schema_version": "office-ir/1.0",
        "kind": "workbook",
        "document_id": "draft:drop",
        "metadata": {"title": "drop"},
        "document": {"worksheets": [{"name": "S", "rows": [["n", 3]]}]},
    }
    compiled = parse_envelope(base)

    richer = parse_envelope(
        {
            **base,
            "document": {
                **base["document"],
                "validations": [
                    {"worksheet": "S", "range": "A2:A9", "kind": "list", "params": {"values": ["a"]}}
                ],
                "conditional_formats": [
                    {"worksheet": "S", "range": "B2:B9", "kind": "cell_value", "params": {"formula": ["5"]}}
                ],
                "tables": [{"worksheet": "S", "range": "A1:B9", "name": "T"}],
            },
        }
    )
    out = compile_workbook(compiled, tmp_path / "dropped.xlsx")
    report = read_back(richer, out)
    assert not report.ok
    dropped = {m.feature for m in report.feature_mismatches}
    assert dropped == {"validation", "conditional_format", "table"}


def test_read_back_detects_dropped_formula(tmp_path: Path) -> None:
    base = {
        "schema_version": "office-ir/1.0",
        "kind": "workbook",
        "document_id": "draft:drop-f",
        "metadata": {"title": "drop-f"},
        "document": {"worksheets": [{"name": "S", "rows": [["a", 2]]}]},
    }
    out = compile_workbook(parse_envelope(base), tmp_path / "nf.xlsx")
    with_formula = parse_envelope(
        {
            **base,
            "document": {
                **base["document"],
                "formulas": [{"worksheet": "S", "cell": "C1", "formula": "=B1*2"}],
            },
        }
    )
    report = read_back(with_formula, out)
    assert not report.ok
    assert report.feature_mismatches[0].feature == "formula"
    assert report.feature_mismatches[0].expected == "=B1*2"


def test_read_back_missing_file_raises(tmp_path: Path) -> None:
    envelope = _envelope(WorkbookIr(worksheets=(Worksheet(name="S"),)))
    with pytest.raises(FileNotFoundError):
        read_back(envelope, tmp_path / "absent.xlsx")


def test_read_back_example_workbook(tmp_path: Path) -> None:
    from dcc_mcp_excel.workbook_ir import load_workbook_ir

    envelope = load_workbook_ir(EXAMPLES / "shot_list.json")
    out = compile_workbook(envelope, tmp_path / "example.xlsx")
    report = read_back(envelope, out)
    assert report.ok, report.to_dict()


@pytest.mark.parametrize(
    "expected,actual,match",
    [
        (None, None, True),
        ("a", "a", True),
        ("a", " a ", True),
        ("a", "b", False),
        (1, 1, True),
        (1, 1.0000000001, True),
        (1, 2, False),
        (True, True, True),
        (True, 1, False),
        (1, True, False),
        ("1", 1, False),
    ],
)
def test_cell_values_match(expected, actual, match) -> None:
    assert cell_values_match(expected, actual) is match
