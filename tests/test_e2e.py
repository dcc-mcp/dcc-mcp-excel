"""End-to-end: production data → XLSX → read-back verification.

The acceptance criterion is the whole chain, not the compile step: take a
production-tracking payload shaped like what `dcc-mcp-fpt` returns
(`find_entities` → `results`), turn it into a Workbook IR, compile it, and
read it back cell by cell. If the adapter silently reordered or dropped a
row, this test fails.
"""

from __future__ import annotations

import json
from pathlib import Path

from dcc_mcp_excel.compiler import compile_workbook
from dcc_mcp_excel.readback import read_back
from dcc_mcp_excel.validate import validate_envelope
from dcc_mcp_excel.workbook_ir import parse_envelope

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"

# A trimmed dcc-mcp-fpt `shotgrid-crud / find_entities` response: the tool
# returns {"success": ..., "message": ..., "context": {"results": [...]}}.
FPT_RESPONSE = {
    "success": True,
    "message": "Found 3 shots",
    "context": {
        "results": [
            {"id": 101, "code": "sh010", "sg_status_list": "ip", "sg_cut_duration": 120, "entity": "Shot"},
            {"id": 102, "code": "sh020", "sg_status_list": "fin", "sg_cut_duration": 84, "entity": "Shot"},
            {"id": 103, "code": "sh030", "sg_status_list": "wip", "sg_cut_duration": 240, "entity": "Shot"},
        ]
    },
}


def fpt_results_to_workbook_ir(response: dict, *, document_id: str = "draft:production-report") -> dict:
    """Map an fpt `find_entities` response onto a Workbook IR envelope.

    Production data is columnar by nature (one row per entity), so the
    mapping is a header row plus one row per entity — the same shape the
    Office production-dashboard skill consumes.
    """
    results = response["context"]["results"]
    headers = ["id", "code", "status", "duration"]
    rows = [[r["id"], r["code"], r["sg_status_list"], r["sg_cut_duration"]] for r in results]
    return {
        "schema_version": "office-ir/1.0",
        "kind": "workbook",
        "document_id": document_id,
        "metadata": {"title": "Production shot report", "author": "dcc-mcp-fpt", "language": "en"},
        "document": {
            "worksheets": [{"name": "Shots", "rows": [headers, *rows]}],
            "tables": [{"worksheet": "Shots", "range": f"A1:D{len(rows) + 1}", "name": "Shots"}],
            "formulas": [{"worksheet": "Shots", "cell": "D5", "formula": f"=SUM(D2:D{len(rows) + 1})"}],
            "calculation_policy": {"mode": "auto", "full_calc_on_load": True},
        },
        "outputs": ["xlsx"],
    }


def test_fpt_production_data_round_trips(tmp_path: Path) -> None:
    ir = fpt_results_to_workbook_ir(FPT_RESPONSE)
    envelope = parse_envelope(ir)

    assert validate_envelope(envelope)["ok"]

    out = compile_workbook(envelope, tmp_path / "production-report.xlsx")
    report = read_back(envelope, out)

    assert report.ok, report.to_dict()
    # header row + one row per entity, four columns each.
    assert report.checked_cells == 4 * (len(FPT_RESPONSE["context"]["results"]) + 1)
    assert report.formulas_present == 1


def test_every_fpt_row_survives(tmp_path: Path) -> None:
    """Row-for-row fidelity: an agent reading the xlsx back sees the same data."""
    from openpyxl import load_workbook

    envelope = parse_envelope(fpt_results_to_workbook_ir(FPT_RESPONSE))
    out = compile_workbook(envelope, tmp_path / "rows.xlsx")
    sheet = load_workbook(str(out))["Shots"]
    rows = list(sheet.iter_rows(min_row=2, max_row=4, values_only=True))
    assert rows == [(101, "sh010", "ip", 120), (102, "sh020", "fin", 84), (103, "sh030", "wip", 240)]


def test_example_workbook_end_to_end(tmp_path: Path) -> None:
    envelope = parse_envelope(json.loads((EXAMPLES / "shot_list.json").read_text(encoding="utf-8")))
    out = compile_workbook(envelope, tmp_path / "shot-list.xlsx")
    report = read_back(envelope, out)
    assert report.ok, report.to_dict()
    assert report.sheets == ("Shots", "Summary")
    assert report.validations == 1
    assert report.conditional_formats == 1
    assert "ShotCount" in report.named_ranges
