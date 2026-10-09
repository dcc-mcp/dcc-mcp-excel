"""Headless compiler tests — Office-free, runs on any platform."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from dcc_mcp_excel.compiler import UnknownWorksheetError, compile_workbook
from dcc_mcp_excel.workbook_ir import (
    Metadata,
    WorkbookEnvelope,
    WorkbookIr,
    Worksheet,
    load_workbook_ir,
    parse_envelope,
)

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"


def _envelope(document: WorkbookIr, document_id: str = "draft:test") -> WorkbookEnvelope:
    return WorkbookEnvelope(
        schema_version="office-ir/1.0",
        kind="workbook",
        document_id=document_id,
        metadata=Metadata(title="Test"),
        document=document,
    )


def test_compile_writes_rows_and_values(tmp_path: Path) -> None:
    envelope = _envelope(
        WorkbookIr(worksheets=(Worksheet(name="Shots", rows=(("shot", "status"), ("sh010", "ip"))),))
    )
    out = compile_workbook(envelope, tmp_path / "out.xlsx")
    assert out.is_file() and out.stat().st_size > 0
    workbook = load_workbook(str(out))
    sheet = workbook["Shots"]
    assert sheet["A1"].value == "shot"
    assert sheet["B2"].value == "ip"


def test_compile_preserves_cell_types(tmp_path: Path) -> None:
    envelope = _envelope(WorkbookIr(worksheets=(Worksheet(name="S", rows=(("text", 7, 2.5, True, None),)),)))
    out = compile_workbook(envelope, tmp_path / "types.xlsx")
    row = next(load_workbook(str(out))["S"].iter_rows(min_row=1, max_row=1, max_col=5, values_only=True))
    assert row == ("text", 7, 2.5, True, None)


def test_compile_adds_suffix_when_missing(tmp_path: Path) -> None:
    envelope = _envelope(WorkbookIr(worksheets=(Worksheet(name="S", rows=(("a",),)),)))
    out = compile_workbook(envelope, tmp_path / "no-suffix")
    assert out.suffix == ".xlsx" and out.is_file()


def test_compile_sanitizes_and_deduplicates_sheet_titles(tmp_path: Path) -> None:
    envelope = _envelope(
        WorkbookIr(
            worksheets=(
                Worksheet(name="Shot/List", rows=(("a",),)),
                Worksheet(name="Shot*List", rows=(("b",),)),
                Worksheet(name="X" * 40, rows=(("c",),)),
            )
        )
    )
    out = compile_workbook(envelope, tmp_path / "titles.xlsx")
    workbook = load_workbook(str(out))
    assert len(workbook.sheetnames) == 3
    assert len(set(workbook.sheetnames)) == 3
    assert all(len(name) <= 31 for name in workbook.sheetnames)
    assert all(not (set(name) & set("\\/*?:[]")) for name in workbook.sheetnames)


def test_compile_writes_formulas(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:f",
            "metadata": {"title": "f"},
            "document": {
                "worksheets": [{"name": "S", "rows": [["a", 2]]}],
                "formulas": [{"worksheet": "S", "cell": "C1", "formula": "=B1*2"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "formula.xlsx")
    assert load_workbook(str(out))["S"]["C1"].value == "=B1*2"


def test_compile_writes_addressed_cells(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:c",
            "metadata": {"title": "c"},
            "document": {
                "worksheets": [{"name": "S", "rows": []}],
                "cells": [{"address": "B2", "value": 42}, {"address": "C2", "formula": "=B2*2"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "cells.xlsx")
    sheet = load_workbook(str(out))["S"]
    assert sheet["B2"].value == 42
    assert sheet["C2"].value == "=B2*2"


def test_unqualified_cell_writes_only_the_first_sheet(tmp_path: Path) -> None:
    """`document.cells` is document-level: one write, onto the first sheet."""
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:multi",
            "metadata": {"title": "multi"},
            "document": {
                "worksheets": [{"name": "A", "rows": []}, {"name": "B", "rows": []}],
                "cells": [{"address": "B2", "value": 42}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "multi.xlsx")
    workbook = load_workbook(str(out))
    assert workbook["A"]["B2"].value == 42
    assert workbook["B"]["B2"].value is None


def test_qualified_cell_targeting_non_first_sheet_compiles(tmp_path: Path) -> None:
    """A sheet qualifier may name any sheet, not just the first one."""
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:multi",
            "metadata": {"title": "multi"},
            "document": {
                "worksheets": [{"name": "A", "rows": []}, {"name": "Summary", "rows": []}],
                "cells": [{"address": "Summary!B2", "value": 7}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "qualified.xlsx")
    workbook = load_workbook(str(out))
    assert workbook["Summary"]["B2"].value == 7
    assert workbook["A"]["B2"].value is None


def test_qualified_cell_resolves_sanitized_sheet_name(tmp_path: Path) -> None:
    """A qualifier names the IR sheet, so a sanitized title stays addressable.

    The IR grammar allows a sheet qualifier only bare (identifier-like) or
    single-quoted, so an illegal title is addressed in its quoted form.
    """
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:multi",
            "metadata": {"title": "multi"},
            "document": {
                "worksheets": [{"name": "A", "rows": []}, {"name": "Sum/mary", "rows": []}],
                "cells": [{"address": "'Sum/mary'!B2", "value": 5}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "sanitized-qualified.xlsx")
    assert load_workbook(str(out))["Sum-mary"]["B2"].value == 5


def test_addressed_cell_naming_unknown_sheet_raises(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:c",
            "metadata": {"title": "c"},
            "document": {
                "worksheets": [{"name": "S", "rows": []}],
                "cells": [{"address": "Ghost!B2", "value": 1}],
            },
        }
    )
    with pytest.raises(KeyError, match="unknown worksheet"):
        compile_workbook(envelope, tmp_path / "ghost.xlsx")


def test_compile_writes_named_ranges(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:n",
            "metadata": {"title": "n"},
            "document": {
                "worksheets": [{"name": "Summary", "rows": [["x", 1]]}],
                "named_ranges": [{"name": "Total", "refers_to": "Summary!$A$1:$B$1"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "named.xlsx")
    assert "Total" in load_workbook(str(out)).defined_names


def test_duplicate_named_range_raises(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:n",
            "metadata": {"title": "n"},
            "document": {
                "worksheets": [{"name": "S", "rows": []}],
                "named_ranges": [
                    {"name": "Dup", "refers_to": "S!$A$1"},
                    {"name": "Dup", "refers_to": "S!$A$2"},
                ],
            },
        }
    )
    with pytest.raises(ValueError, match="duplicate named range"):
        compile_workbook(envelope, tmp_path / "dup.xlsx")


def test_compile_writes_list_validation(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:v",
            "metadata": {"title": "v"},
            "document": {
                "worksheets": [{"name": "S", "rows": [["status"]]}],
                "validations": [
                    {"worksheet": "S", "range": "A2:A99", "kind": "list", "params": {"values": ["ip", "fin"]}}
                ],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "validation.xlsx")
    validations = load_workbook(str(out))["S"].data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].formula1 == '"ip,fin"'


def test_list_validation_without_values_raises(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:v",
            "metadata": {"title": "v"},
            "document": {
                "worksheets": [{"name": "S", "rows": []}],
                "validations": [{"worksheet": "S", "range": "A2:A9", "kind": "list", "params": {}}],
            },
        }
    )
    with pytest.raises(KeyError, match=r"params\.values"):
        compile_workbook(envelope, tmp_path / "v.xlsx")


@pytest.mark.parametrize(
    "kind,params",
    [
        ("cell_value", {"operator": "greaterThan", "formula": ["5"]}),
        ("color_scale", {"colors": ["F8696B", "63BE7B"]}),
        ("data_bar", {"start_value": 0, "end_value": 10}),
        ("formula", {"formula": ["=A1>5"]}),
    ],
)
def test_compile_writes_conditional_formats(tmp_path: Path, kind: str, params: dict) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:cf",
            "metadata": {"title": "cf"},
            "document": {
                "worksheets": [{"name": "S", "rows": [["n", 3]]}],
                "conditional_formats": [{"worksheet": "S", "range": "B2:B9", "kind": kind, "params": params}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / f"cf-{kind}.xlsx")
    # _cf_rules maps a ConditionalFormattingRange (carrying .sqref) to its rules.
    rules = load_workbook(str(out))["S"].conditional_formatting._cf_rules
    assert rules
    applied = [entry for entry in rules if str(entry.sqref) == "B2:B9"]
    assert len(applied) == 1
    assert len(rules[applied[0]]) == 1


def test_compile_raises_on_unsupported_conditional_format_params(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:cf",
            "metadata": {"title": "cf"},
            "document": {
                "worksheets": [{"name": "S", "rows": []}],
                "conditional_formats": [{"worksheet": "S", "range": "B2:B9", "kind": "formula", "params": {}}],
            },
        }
    )
    with pytest.raises(ValueError, match=r"params\.formula"):
        compile_workbook(envelope, tmp_path / "cf-bad.xlsx")


def test_compile_applies_calculation_policy_metadata(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:calc",
            "metadata": {"title": "calc", "author": "pipeline"},
            "document": {
                "worksheets": [{"name": "S", "rows": [["a", 1]]}],
                "calculation_policy": {"mode": "auto", "full_calc_on_load": True},
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "calc.xlsx")
    workbook = load_workbook(str(out))
    assert workbook.properties.title == "calc"
    assert workbook.calculation.fullCalcOnLoad is True


def test_compile_writes_validation_on_sanitized_sheet_title(tmp_path: Path) -> None:
    """A title needing sanitization must not silently drop its validation.

    `Shot/List` compiles to `Shot-List`. Matching the spec against the
    compiled title skipped it; the IR name is the only correct key.
    """
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:sanitized",
            "metadata": {"title": "sanitized"},
            "document": {
                "worksheets": [{"name": "Shot/List", "rows": [["status"]]}],
                "validations": [
                    {"worksheet": "Shot/List", "range": "A2:A99", "kind": "list", "params": {"values": ["ip", "fin"]}}
                ],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "sanitized.xlsx")
    sheet = load_workbook(str(out))["Shot-List"]
    assert len(sheet.data_validations.dataValidation) == 1


def test_compile_writes_conditional_format_on_sanitized_sheet_title(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:sanitized",
            "metadata": {"title": "sanitized"},
            "document": {
                "worksheets": [{"name": "Shot/List", "rows": [["n", 3]]}],
                "conditional_formats": [
                    {"worksheet": "Shot/List", "range": "B2:B9", "kind": "cell_value", "params": {"formula": ["5"]}}
                ],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "cf-sanitized.xlsx")
    rules = load_workbook(str(out))["Shot-List"].conditional_formatting._cf_rules
    applied = [entry for entry in rules if str(entry.sqref) == "B2:B9"]
    assert len(applied) == 1


def test_compile_writes_formula_on_sanitized_sheet_title(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:sanitized",
            "metadata": {"title": "sanitized"},
            "document": {
                "worksheets": [{"name": "Shot/List", "rows": [["a", 2]]}],
                "formulas": [{"worksheet": "Shot/List", "cell": "C1", "formula": "=B1*2"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "f-sanitized.xlsx")
    assert load_workbook(str(out))["Shot-List"]["C1"].value == "=B1*2"


@pytest.mark.parametrize(
    "feature_key,spec",
    [
        ("formulas", {"cell": "C1", "formula": "=B1*2"}),
        ("validations", {"range": "A2:A9", "kind": "list", "params": {"values": ["a"]}}),
        ("conditional_formats", {"range": "B2:B9", "kind": "cell_value", "params": {"formula": ["5"]}}),
        ("tables", {"range": "A1:B9", "name": "T"}),
    ],
)
def test_compile_raises_when_feature_names_unknown_worksheet(
    tmp_path: Path, feature_key: str, spec: dict
) -> None:
    """An unwritable spec is an error, never a silently dropped feature."""
    document = {"worksheets": [{"name": "S", "rows": []}], feature_key: [{"worksheet": "Ghost", **spec}]}
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:ghost",
            "metadata": {"title": "ghost"},
            "document": document,
        }
    )
    with pytest.raises(UnknownWorksheetError, match="unknown worksheet 'Ghost'"):
        compile_workbook(envelope, tmp_path / "ghost.xlsx")


def test_compile_writes_tables(tmp_path: Path) -> None:
    envelope = parse_envelope(
        {
            "schema_version": "office-ir/1.0",
            "kind": "workbook",
            "document_id": "draft:t",
            "metadata": {"title": "t"},
            "document": {
                "worksheets": [{"name": "Shots", "rows": [["shot", "status"], ["sh010", "ip"]]}],
                "tables": [{"worksheet": "Shots", "range": "A1:B2", "name": "ShotTable"}],
            },
        }
    )
    out = compile_workbook(envelope, tmp_path / "tables.xlsx")
    assert list(load_workbook(str(out))["Shots"].tables) == ["ShotTable"]


def test_compile_example_workbook(tmp_path: Path) -> None:
    envelope = load_workbook_ir(EXAMPLES / "shot_list.json")
    out = compile_workbook(envelope, tmp_path / "example.xlsx")
    workbook = load_workbook(str(out))
    assert workbook.sheetnames == [ws.name for ws in envelope.document.worksheets]
    assert out.stat().st_size > 0
